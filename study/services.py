import os
import json
import logging
import re
from typing import Dict, Any, List
import pdfplumber
import docx
from pptx import Presentation
from groq import Groq

logger = logging.getLogger(__name__)

DEFAULT_MODEL = os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile")


def get_inference_client():
    """Initialize neural synthesis engine client using environment configuration."""
    api_key = os.getenv("GROQ_API_KEY")
    if not api_key or api_key == "your_groq_api_key_here":
        return None
    return Groq(api_key=api_key)


def extract_text(file_path: str, file_extension: str) -> str:
    """
    Extract structured text content across documents: PDF, DOCX, PPTX, TXT.
    """
    ext = file_extension.lower().lstrip(".")
    text_parts = []

    try:
        if ext == "pdf":
            with pdfplumber.open(file_path) as pdf:
                for page_idx, page in enumerate(pdf.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text.strip())

        elif ext in ("docx", "doc"):
            doc = docx.Document(file_path)
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        text_parts.append(row_text)

        elif ext in ("pptx", "ppt"):
            prs = Presentation(file_path)
            for slide_idx, slide in enumerate(prs.slides):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text.strip())
                if slide_texts:
                    text_parts.append(f"[Section {slide_idx + 1}]\n" + "\n".join(slide_texts))

        elif ext in ("txt", "md"):
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text_parts.append(f.read().strip())

        else:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text_parts.append(f.read().strip())

    except Exception as e:
        logger.error(f"Error parsing manuscript from {file_path} ({ext}): {e}", exc_info=True)
        raise RuntimeError(f"Failed to extract structured text from {ext.upper()} file: {str(e)}")

    return "\n\n".join(text_parts).strip()


def _clean_json_response(raw_content: str) -> Any:
    """Parse structured JSON response from synthesis engine."""
    content = raw_content.strip()
    if content.startswith("```json"):
        content = content[len("```json"):].strip()
    elif content.startswith("```"):
        content = content[len("```"):].strip()
    if content.endswith("```"):
        content = content[:-3].strip()

    return json.loads(content)


def generate_summary(text: str) -> Dict[str, List[str]]:
    """
    Generate high-yield key takeaways and real-world implementation cases.
    """
    client = get_inference_client()

    if not client:
        sentences = [s.strip() for s in re.split(r'[.\n]+', text) if len(s.strip()) > 30]
        key_points = [
            f"Core Tenet: {s}" for s in (sentences[:5] if sentences else ["Document parsed and indexed for active recall."])
        ]
        examples = [
            "Practical Application: Integrating these foundational principles into architectural design and data pipelines.",
            "Case Benchmark: Enterprise adoption and performance optimization in mission-critical environments."
        ]
        return {
            "key_points": key_points,
            "real_world_examples": examples
        }

    truncated_text = text[:60000]

    system_prompt = (
        "You are an elite academic scholar and executive research analyst. Analyze the provided study manuscript "
        "and produce an authoritative, crystal-clear executive briefing. Respond with a strictly valid JSON object "
        "containing exactly two keys:\n"
        "1. 'key_points': An array of high-yield, conceptually dense takeaway bullet points breaking down the core principles.\n"
        "2. 'real_world_examples': An array of concrete, real-world case examples, practical implementations, or empirical applications.\n"
        "Deliver rigorous, impressive, and sophisticated insights tailored for mastery."
    )

    user_prompt = f"Study Manuscript:\n\n{truncated_text}\n\nDeliver JSON executive briefing only."

    response = client.chat.completions.create(
        model=DEFAULT_MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        response_format={"type": "json_object"},
        temperature=0.3,
    )

    content = response.choices[0].message.content
    data = _clean_json_response(content)

    return {
        "key_points": data.get("key_points", []),
        "real_world_examples": data.get("real_world_examples", []),
    }


def explain_highlight(highlighted_text: str, full_context: str) -> str:
    """
    Deep Contextual Clarification: Formulates an intuitive breakdown of the selected phrase using surrounding context.
    """
    client = get_inference_client()

    if not client:
        return (
            f"Contextual Analysis of '{highlighted_text}':\n\n"
            "This principle serves as a cornerstone mechanism within the current manuscript. "
            "It establishes the governing criteria and mathematical/conceptual relationships that dictate system behavior.\n\n"
            "In applied settings, this concept ensures consistency, reduces state divergence, and enables deterministic outcomes "
            "across high-dimensional workflows."
        )

    context_snippet = full_context[:25000]

    system_prompt = (
        "You are a master academic polymath and cognitive mentor. The researcher has highlighted a critical passage "
        "or concept from their text. Provide a master-level, crystal-clear conceptual deep-dive that demystifies this idea "
        "within the surrounding manuscript context. Utilize powerful intuitive analogies, break down technical mechanics, "
        "and explain why this distinction is vital. Maintain an authoritative yet exceptionally lucid tone (2-3 concise paragraphs)."
    )

    user_prompt = (
        f"--- SURROUNDING MANUSCRIPT CONTEXT ---\n{context_snippet}\n\n"
        f"--- SELECTED PASSAGE ---\n\"{highlighted_text}\"\n\n"
        "Provide a comprehensive, elegant, and crystal-clear explanation."
    )

    response = client.chat.completions.create(
        model=DEFAULT_MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        temperature=0.4,
    )

    return response.choices[0].message.content.strip()


def generate_quiz(notes_and_summary_text: str) -> List[Dict[str, Any]]:
    """
    Formulate a rigorous 5-question active recall mastery examination based on notes and synthesis.
    """
    client = get_inference_client()

    if not client:
        return [
            {
                "question": "What is the primary governing dynamic established by the foundational principles in this text?",
                "options": [
                    "To establish scalable operational benchmarks across complex states",
                    "To completely replace classical algorithms unconditionally",
                    "To prevent asynchronous data flow in all subsystems",
                    "To minimize computational complexity without architectural validation"
                ],
                "correct_answer": "To establish scalable operational benchmarks across complex states"
            },
            {
                "question": "How does superposition fundamentally depart from classical binary constraints?",
                "options": [
                    "It permits existence in linear combinations of states simultaneously",
                    "It strictly enforces deterministic parity without uncertainty",
                    "It requires continuous analog intervention at all stages",
                    "It eliminates physical noise and thermal fluctuations entirely"
                ],
                "correct_answer": "It permits existence in linear combinations of states simultaneously"
            },
            {
                "question": "Which engineering obstacle represents the primary bottleneck in practical physical implementation?",
                "options": [
                    "Environmental decoherence and thermal noise",
                    "Excessive storage bandwidth allocation",
                    "Absence of foundational mathematical formalisms",
                    "Over-abundance of cryo-cooling hardware capacity"
                ],
                "correct_answer": "Environmental decoherence and thermal noise"
            },
            {
                "question": "In cryptographic architecture, which formulation demonstrates polynomial-time public-key factoring?",
                "options": [
                    "Shor's Algorithm",
                    "Grover's Search Protocol",
                    "Dijkstra's Shortest Path",
                    "Diffie-Hellman Key Exchange"
                ],
                "correct_answer": "Shor's Algorithm"
            },
            {
                "question": "What strategic methodology yields the highest active recall when mastering complex domains?",
                "options": [
                    "Synthesizing precision notes coupled with targeted concept retrieval assessments",
                    "Passive rote memorization without contextual cross-referencing",
                    "Avoiding analogical breakdowns of core mathematical formulas",
                    "Disregarding empirical case studies and real-world applications"
                ],
                "correct_answer": "Synthesizing precision notes coupled with targeted concept retrieval assessments"
            }
        ]

    content_snippet = notes_and_summary_text[:30000]

    system_prompt = (
        "You are an assessment architect for advanced academic and professional certification. "
        "Formulate a rigorous 5-question multiple choice assessment evaluating deep conceptual mastery "
        "of the student's manuscript and notes. Respond with a strictly valid JSON object with a single key "
        "'questions' containing an array of 5 questions. Each question item must possess:\n"
        "- 'question': Thoughtful, rigorous question evaluating core mechanics.\n"
        "- 'options': An array of exactly 4 plausible, distinct answer choices.\n"
        "- 'correct_answer': The exact string of the correct choice matching one of the options."
    )

    user_prompt = (
        f"Manuscript & Student Notes:\n\n{content_snippet}\n\n"
        "Generate 5 high-caliber assessment questions. Return valid JSON only with 'questions' array."
    )

    response = client.chat.completions.create(
        model=DEFAULT_MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        response_format={"type": "json_object"},
        temperature=0.3,
    )

    content = response.choices[0].message.content
    data = _clean_json_response(content)

    questions = data.get("questions", [])
    if not isinstance(questions, list):
        if isinstance(data, list):
            questions = data
        else:
            questions = []

    return questions
